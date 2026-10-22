"""Tests de l'onglet Textile, de la règle fournisseur, des graphes SVG et de l'export Excel (données FICTIVES)."""

import io
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from core import charts, xlsx
from core import io as sio
from core import textile as tx
from core import transforms as tf

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests" / "fixtures"
TODAY = pd.Timestamp("2026-10-22")


def octets(nom):
    return (FIX / nom).read_bytes()


@pytest.fixture(scope="module")
def tables():
    stock, _ = sio.lire_etat_stock(octets("etat_stock.xlsx"))
    cmd, _ = sio.lire_commandes_sap(octets("commandes_sap.xlsx"))
    ref, feuille = sio.lire_references_textile(octets("references_textile.xlsx"))
    assert feuille == "PN Catalog"
    return tx.construire_tables(stock, cmd, ref, TODAY, True)


# --------------------------------------------------------------------------
# Règle fournisseur (demande métier : jamais « fournisseur inconnu »)
# --------------------------------------------------------------------------
def test_fournisseur_retrouve_par_le_code():
    lignes = pd.DataFrame({
        "PN": ["A", "A", "B", "C", "D", "D"],
        "Code_fournisseur_cmd": ["F1", "F1", "F2", np.nan, "F9", np.nan],
        "Fournisseur": ["Alpha", np.nan, np.nan, np.nan, np.nan, np.nan],
    })
    lignes.loc[6] = ["B", "F2", "Bravo"]
    out = tf.completer_fournisseurs(lignes).tolist()
    assert out[:3] == ["Alpha", "Alpha", "Bravo"]      # même code ailleurs dans le fichier
    assert pd.isna(out[3])                             # ni nom ni code, sur aucune ligne du PN
    assert out[4] == "Code F9" and out[5] == "Code F9" # code seul, puis code du même PN


def test_fournisseur_complete_dans_le_fichier_commandes():
    cmd, _ = sio.lire_commandes_sap(octets("commandes_sap.xlsx"))
    ligne = cmd[cmd["Commande"] == "CMD-0031"].iloc[0]
    assert ligne["Fournisseur"] == "Fournisseur Charlie"   # vide dans le fichier, code F003


def test_pn_sans_espaces():
    assert sio.cle_pn(pd.Series([" ab 12 ", "X Y"])).tolist() == ["AB12", "XY"]


# --------------------------------------------------------------------------
# Textile
# --------------------------------------------------------------------------
def test_tables_textile(tables):
    s, o = tables
    assert set(s["PNR"]) == {"FX-1001", "FX-1002", "FX-1003", "FX-1004", "FX-1005", "FX-1006", "FX-1007", "FX-1012",
                             "FX-1015", "FX-2001", "FX-2002"}          # FX-9999 : au catalogue mais pas au stock
    assert tx.annees_conso(s) == [2021, 2022, 2023, 2024, 2025]
    st = s.set_index("PNR")["Statut activité"]
    assert st["FX-2001"] == "Dormant" and st["FX-2002"] == "Inactif" and st["FX-1003"] == "Actif"
    assert s.set_index("PNR").loc["FX-2001", "Valeur stock dormant (€)"] == 720
    assert "Evolution 2025 vs 2024 (%)" in s.columns
    assert s.set_index("PNR").loc["FX-1015", "Commandes (liste)"] == "CMD-0151 ; CMD-0152"


def test_statuts_commandes(tables):
    _, o = tables
    st = o.set_index("Commande")["Statut commande"]
    assert st["CMD-0031"] == "En attente BFE"          # date promise 2049
    assert st["CMD-0007"] == "Sans date promise"       # 2031 : date fictive
    assert st["CMD-0006"] == "Sans date promise"       # vide
    assert st["CMD-0003"] == "A l'heure"               # promise 10/11 vs demandée 01/11 : 9 j ≤ 10 j de tolérance
    assert st["CMD-0005"] == "En retard"
    assert o.set_index("Commande").loc["CMD-0005", "Classe de retard"] == "51 à 150 j"


def test_relances_ordre(tables):
    _, o = tables
    rel = tx.relances(o)
    statuts = rel["Statut commande"].tolist()
    assert statuts == sorted(statuts, key=["En retard", "Sans date promise", "En attente BFE"].index)
    retards = rel.loc[rel["Statut commande"] == "En retard", "Jours de retard"].tolist()
    assert retards == sorted(retards, reverse=True)
    assert "CMD-0012" not in set(rel["Commande"])      # livrée (reste 0) : pas de relance


def test_syntheses(tables):
    s, o = tables
    r = tx.synthese(s, o)
    assert r["pn"] == 11 and r["dormants"] == 1 and r["inactifs"] == 1
    assert tx.resume_categories(s, o)["Catégorie"].tolist() == ["CEINTURE", "HOUSSE", "MOUSSE", "RIDEAUX", "TOTAL"]
    poids = tx.poids_par_categorie(s, o)
    assert poids.loc["PN"].sum() == 11
    assert tx.tableau_retards(o)["Lignes"].sum() == (o["Statut commande"] == "En retard").sum()
    assert not tx.planning_livraisons(o, TODAY).empty


def test_filtre_catalogue_identique():
    """Le filtre en amont sur le catalogue donne le même résultat que le filtre final d'origine."""
    stock, _ = sio.lire_etat_stock(octets("etat_stock.xlsx"))
    cmd, _ = sio.lire_commandes_sap(octets("commandes_sap.xlsx"))
    ref, _ = sio.lire_references_textile(octets("references_textile.xlsx"))
    s1, o1 = tx.construire_tables(stock, cmd, ref, TODAY, True)
    s2, o2 = tx.construire_tables(stock, cmd, ref, TODAY, False)
    s2 = s2[s2["REF_Equipement"].notna()].reset_index(drop=True)
    o2 = o2[o2["REF_Equipement"].notna()].reset_index(drop=True)
    pd.testing.assert_frame_equal(s1, s2[s1.columns])
    pd.testing.assert_frame_equal(o1, o2)


# --------------------------------------------------------------------------
# Agrégation de textes : même résultat que la version d'origine (groupby + fonction)
# --------------------------------------------------------------------------
def test_joindre_par_groupe_equivalent():
    rng = np.random.default_rng(0)
    d = pd.DataFrame({"k": rng.choice(list("abcdefg"), 400), "v": rng.choice(["x", "y", "z", " ", None, "nan", "w"], 400)})
    attendu = d.groupby("k")["v"].agg(tf.joindre_uniques).dropna()
    obtenu = tf.joindre_par_groupe(d, "k", "v").reindex(attendu.index)
    assert obtenu.tolist() == attendu.tolist()


# --------------------------------------------------------------------------
# Graphes
# --------------------------------------------------------------------------
def test_graphes_html():
    h = charts.barres(["A", "B"], {"x": [3, np.nan]}, ligne={"l": [10, 20]}, fmt_ligne=charts.pourcent)
    assert "<svg" in h and h.count("<rect") == 2 and "20 %" in h
    assert "s360-rotpad" in charts.barres([f"Fournisseur {i}" for i in range(9)], {"x": list(range(9))})
    assert "<svg" not in charts.parts(pd.DataFrame({"a": [1, 3]}, index=["L1", "L2"]), {"a": "#000"})
    assert charts.vide() in charts.parts(pd.DataFrame(), {})
    assert charts.euros(1_234_567) == "1,2 M€" and charts.court(12_345) == "12k" and charts.entier(1234) == "1 234"


# --------------------------------------------------------------------------
# Export Excel (écrivain maison, bibliothèque standard)
# --------------------------------------------------------------------------
def _classeur_test():
    df = pd.DataFrame({"PN": ["A1", "B<&>", None], "Qté": [1, 2.5, np.nan],
                       "Date": pd.to_datetime(["2026-11-15", None, "2027-01-02"]), "Texte": ["é — ü", "", "x"]})
    return xlsx.classeur({"Synthèse": pd.DataFrame({"Indicateur": ["a"], "Valeur": [3]}), "Détail": df},
                         {"Détail": ["F8D7DA", None, "D9F2E6"]}), df


def test_xlsx_relu_par_openpyxl():
    data, df = _classeur_test()
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Détail"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D4"
    assert ws["A1"].font.b and ws["A1"].fill.fgColor.rgb.endswith("002157")
    assert ws["C2"].number_format.lower() == "dd/mm/yyyy" and ws["C2"].value.date().isoformat() == "2026-11-15"
    assert ws["A3"].value == "B<&>" and ws["D2"].value == "é — ü" and ws["B3"].value == 2.5
    assert ws["A2"].fill.fgColor.rgb.endswith("F8D7DA")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice absent")
def test_xlsx_ouvert_par_libreoffice(tmp_path):
    data, _ = _classeur_test()
    f = tmp_path / "test.xlsx"
    f.write_bytes(data)
    subprocess.run(["soffice", "--headless", "--convert-to", "csv", "--outdir", str(tmp_path), str(f)],
                   check=True, capture_output=True, timeout=120)
    csv = (tmp_path / "test.csv").read_text(encoding="utf-8", errors="replace")
    assert "Indicateur" in csv  # première feuille relue par un vrai tableur
