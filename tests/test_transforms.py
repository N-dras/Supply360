"""Tests des lectures et croisements (sans Streamlit), sur données FICTIVES."""

import ast
import io
import sys
import types
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from core import io as sio
from core import transforms as tf

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests" / "fixtures"
sys.path.insert(0, str(FIX))
from generer_fixtures import CAS_ATTENDUS, DATE_CHANTIER, GESTIONS  # noqa: E402


def octets(nom: str) -> bytes:
    return (FIX / nom).read_bytes()


@pytest.fixture(scope="module")
def sources():
    besoin = tf.agreger_besoins(sio.lire_besoins_geav(octets("besoins_chantier.xlsx"))[0])
    stock = tf.agreger_stock(sio.lire_etat_stock(octets("etat_stock.xlsx"))[0])
    commandes = tf.agreger_commandes(sio.lire_commandes_sap(octets("commandes_sap.xlsx"))[0])
    return besoin, stock, commandes


@pytest.fixture(scope="module")
def croise(sources):
    return tf.croiser_geav(*sources, GESTIONS, pd.Timestamp(DATE_CHANTIER))


# --------------------------------------------------------------------------
# Lecture / normalisation
# --------------------------------------------------------------------------
def test_normaliser():
    assert sio.normaliser("  QTE SUR\nÉTAGÈRE ") == "qte_sur_etagere"
    assert sio.normaliser("Date de livraison demandée") == "date_de_livraison_demandee"


def test_trouver_colonne_exacte_puis_partielle():
    df = pd.DataFrame(columns=["Code fournisseur", "Fournisseur principal", "PN"])
    assert sio.trouver_colonne(df, "pn", exact=True) == "PN"
    assert sio.trouver_colonne(df, "fournisseur", exact=True) is None
    assert sio.trouver_colonne(df, "fournisseur") == "Code fournisseur"


def test_cle_pn_texte():
    s = pd.Series([123456.0, "  ab-12 ", "000789", None, "", 42])
    assert sio.cle_pn(s).tolist()[:3] == ["123456", "AB-12", "000789"]
    assert sio.cle_pn(s).isna().tolist() == [False, False, False, True, True, False]
    assert sio.cle_pn(s).iloc[5] == "42"


def test_feuilles_et_entetes_detectees():
    assert sio.lire_besoins_geav(octets("besoins_chantier.xlsx"))[1] == "Workpages_PartRequirements"
    assert sio.lire_etat_stock(octets("etat_stock.xlsx"))[1] == "Export"
    assert sio.lire_commandes_sap(octets("commandes_sap.xlsx"))[1] == "Toutes les commandes"


def test_colonne_manquante_message_lisible():
    wb = Workbook()
    wb.active.append(["MPN Prime", "Description"])
    wb.active.append(["FX-1", "x"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(sio.ColonneIntrouvable, match="Colonne « Qty » introuvable dans le fichier Besoins chantier"):
        sio.lire_besoins_geav(buf.getvalue())


def test_fichier_non_excel():
    with pytest.raises(sio.FichierIllisible):
        sio.lire_etat_stock(b"pas un classeur")


def test_asreq():
    lignes, feuille = sio.lire_besoins_asreq(octets("etude_asreq.xlsm"))
    besoin = tf.agreger_besoins(lignes).set_index("PN")
    assert feuille == "Ordered"
    assert besoin.loc["FX-1002", "QTY_GEAV"] == 2
    assert besoin.loc["FX-1002", "TRFX"] == "BC-0002 | BC-0003"  # deux lignes, besoins additionnés (1 + 1)
    assert "Description" not in besoin.columns


# --------------------------------------------------------------------------
# Agrégations
# --------------------------------------------------------------------------
def test_besoins_sommes_par_pn(sources):
    besoin = sources[0].set_index("PN")
    assert besoin.loc["FX-1002", "QTY_GEAV"] == 5
    assert besoin.loc["FX-1002", "Nb_lignes"] == 2
    # deux TRFX (et deux RRS) pour un même PN : un seul besoin, quantités additionnées (2 + 3)
    assert besoin.loc["FX-1002", "TRFX"] == "TRFX-02 | TRFX-18" and besoin.loc["FX-1002", "RRS"] == "RRS-0002 | RRS-0003"
    assert "FX-1013" in besoin.index  # « fx-1013 » normalisé
    assert len(besoin) == 16  # la ligne sans PN est ignorée


def test_commandes_psd_la_plus_proche(sources):
    cmd = sources[2].set_index("PN")
    assert cmd.loc["FX-1015", "PSD"] == pd.Timestamp("2026-11-25")
    assert cmd.loc["FX-1015", "Commande"] == "CMD-0152 | CMD-0151"
    assert cmd.loc["FX-1015", "Nb_commandes"] == 2
    assert bool(cmd.loc["FX-1006", "PSD_est_estimee"]) is True
    assert bool(cmd.loc["FX-1012", "Has_risk_class"]) is True


def test_stock_classe_risque_prioritaire():
    lignes = pd.DataFrame({"PN": ["A", "A", "B", "B", "C"], "CLASS_CODE_STOCK": [1.0, 9.0, np.nan, 2.0, np.nan]})
    r = tf._premier_prioritaire(lignes, "PN", "CLASS_CODE_STOCK", tf.CLASSES_RISQUE)
    assert r["A"] == 9.0 and r["B"] == 2.0 and np.isnan(r["C"])


# --------------------------------------------------------------------------
# Croisement
# --------------------------------------------------------------------------
@pytest.mark.parametrize("pn,code", sorted(CAS_ATTENDUS.items()))
def test_action_par_cas(croise, pn, code):
    assert croise.set_index("PN").loc[pn, "ACTION_CODE"] == code


def test_quantites(croise):
    m = croise.set_index("PN")
    assert (m.loc["FX-1003", "Deficit"], m.loc["FX-1003", "Qty_a_commander"]) == (4, 2)
    # contact = Code agent (CODE AGT APPRO de l'État du stock), plus le code acheteur des commandes
    assert "Commander le complément (2) + relancer Fournisseur Charlie (contact: AG02)" in m.loc["FX-1003", "ACTION"]
    assert m.loc["FX-1008", "ACTION"].startswith("🟠 Relancer Fournisseur Hotel")


def test_gestion_commandes_fait_foi(croise):
    m = croise.set_index("PN")
    assert m.loc["FX-1014", "GESTION"] == 245 and bool(m.loc["FX-1014", "gestion_in_scope"])


def test_classe_risque_prefixe(croise):
    m = croise.set_index("PN")
    assert m.loc["FX-1012", "ACTION"].startswith("🚨 CLASSE 8/9 — PRIORITÉ : 🔴 Commander 2")
    assert not m.loc["FX-1010", "ACTION"].startswith("🚨")


def test_synthese_et_priorisation(croise):
    scope = croise[croise["gestion_in_scope"]]
    s = tf.synthese_geav(scope)
    assert s["n_total"] == 14
    assert (s["n_commander"], s["qty_a_commander"]) == (3, 8)
    assert (s["n_relance"], s["n_relance_urgente"]) == (3, 1)
    actions = tf.actions_a_lancer(scope)
    assert actions["PN"].iloc[0] == "FX-1012"  # classe 9 en tête
    assert actions["ACTION_CODE"].tolist()[1:] == ["commander", "commander", "relancer_urgent",
                                                    "relancer", "relancer", "inconnu"]


def test_jours_avant():
    assert tf.jours_avant(date(2026, 12, 1), date(2026, 11, 1)) == 30


# --------------------------------------------------------------------------
# Non-régression : même résultat que le script d'origine (legacy/streamlit_GEAV.py)
# --------------------------------------------------------------------------
def _legacy():
    """Charge les fonctions du script d'origine sans exécuter son interface Streamlit."""
    source = (ROOT / "legacy" / "streamlit_GEAV.py").read_text(encoding="utf-8")
    arbre = ast.parse(source)
    arbre.body = [n for n in arbre.body if isinstance(n, (ast.FunctionDef, ast.Import, ast.ImportFrom))
                  and not (isinstance(n, ast.Import) and n.names[0].name == "streamlit")
                  and not (isinstance(n, ast.ImportFrom) and n.module and "plotly" in n.module)
                  and not (isinstance(n, ast.Import) and "plotly" in n.names[0].name)]
    st = types.SimpleNamespace(cache_data=lambda **kw: (lambda f: f))
    ns = {"st": st}
    exec(compile(arbre, "legacy", "exec"), ns)
    return ns


def _comparer(attendu: pd.DataFrame, obtenu: pd.DataFrame):
    a = attendu.set_index("PN").sort_index()
    o = obtenu.set_index("PN").sort_index()
    a_msg = (a["ACTION"].str.replace("fichier Stock", "fichier État du stock")
             .str.replace("fournisseur inconnu", tf.FOURNISSEUR_ABSENT))  # libellé modifié à la demande métier
    pd.testing.assert_series_equal(a["ACTION_CODE"], o["ACTION_CODE"], check_dtype=False)
    pd.testing.assert_series_equal(a_msg, o["ACTION"], check_dtype=False, check_names=False)
    for col in ["Deficit", "Qty_a_commander", "QTY_STOCK", "QTY_CDE", "GESTION", "RISK_CLASS"]:
        pd.testing.assert_series_equal(a[col], o[col], check_dtype=False, check_names=False)


def test_identique_au_script_origine_sur_fixtures(croise):
    lg = _legacy()
    besoin, _ = lg["load_besoin"](octets("besoins_chantier.xlsx"))
    stock, _ = lg["load_stock"](octets("etat_stock.xlsx"))
    cmd, _ = lg["load_commandes"](octets("commandes_sap.xlsx"))
    # Seul écart voulu avec le script d'origine : le contact est le CODE AGT APPRO (le script ne l'utilisait qu'à
    # défaut de code acheteur) ; on lui donne donc des commandes sans code acheteur.
    cmd = cmd.assign(Code_acheteur=np.nan)
    attendu = lg["build_cross"](besoin, stock, cmd, GESTIONS, pd.Timestamp(DATE_CHANTIER))
    _comparer(attendu, croise)


@pytest.mark.parametrize("graine", range(5))
def test_identique_au_script_origine_aleatoire(graine):
    """Croisement sur des données aléatoires : les deux implémentations doivent concorder."""
    rng = np.random.default_rng(graine)
    n = 300
    pns = [f"P{i:04d}" for i in range(n)]
    besoin = pd.DataFrame({"PN": pns, "QTY_GEAV": rng.integers(0, 6, n), "Nb_lignes": 1,
                           "Description": "x", "TRFX": rng.choice(["T1", None], n)})
    dans_stock = rng.random(n) < 0.85
    stock = pd.DataFrame({
        "PN": np.array(pns)[dans_stock],
        "QTY_STOCK": rng.integers(0, 6, dans_stock.sum()).astype(float),
        "QTY_CDE": rng.integers(0, 4, dans_stock.sum()).astype(float),
        "GESTION": rng.choice([245.0, 300.0, np.nan], dans_stock.sum()),
        "CODE_AGENT": rng.choice(["AG1", None], dans_stock.sum()),
        "CODE_FOURNISSEUR": rng.choice(["F1", None], dans_stock.sum()),
        "DERNIER_FOURNISSEUR": rng.choice(["Dernier", None], dans_stock.sum()),
        "Description_stock": "y",
        "CCN": rng.choice(["1", None], dans_stock.sum()),
        "CLASS_CODE_STOCK": rng.choice([1.0, 8.0, np.nan], dans_stock.sum()),
    })
    a_cmd = rng.random(n) < 0.6
    k = a_cmd.sum()
    psd = pd.Series(pd.to_datetime("2026-11-01") + pd.to_timedelta(rng.integers(0, 120, k), "D"))
    psd[rng.random(k) < 0.15] = pd.Timestamp("2031-12-31")
    psd[rng.random(k) < 0.15] = pd.NaT
    commandes = pd.DataFrame({
        "PN": np.array(pns)[a_cmd], "Commande": "C", "Code_acheteur": rng.choice(["AC", None], k),
        "Gestion_cmd": rng.choice([245.0, 246.0, 300.0, np.nan], k),
        "Code_fournisseur_cmd": rng.choice(["FC", None], k), "Fournisseur": rng.choice(["Fourn", None], k),
        "PSD": psd.values, "PSD_est_estimee": False, "Statut": "En cours", "Qty_restant_cmd": 1.0,
        "Nb_commandes": 1, "Class_code": rng.choice([1.0, 9.0, np.nan], k),
        "Has_risk_class": rng.choice([True, False], k), "CCN_cmd": "2",
    })
    date_chantier = pd.Timestamp("2026-12-15")
    attendu = _legacy()["build_cross"](besoin, stock, commandes.assign(Code_acheteur=None), [245, 246], date_chantier)
    obtenu = tf.croiser_geav(besoin, stock, commandes, [245, 246], date_chantier)
    assert set(obtenu["ACTION_CODE"]) >= {"ok", "commander", "relancer", "relancer_urgent", "en_cours",
                                          "hors_perimetre", "inconnu"}
    _comparer(attendu, obtenu)


# --------------------------------------------------------------------------
# Export Excel
# --------------------------------------------------------------------------
def test_export_excel_mis_en_forme(croise):
    df = croise[["PN", "QTY_GEAV", "PSD", "ACTION_CODE"]]
    data = sio.exporter_excel({"Synthèse": pd.DataFrame({"Indicateur": ["a"], "Valeur": [1]}), "Détail": df},
                              couleurs_lignes={"Détail": df["ACTION_CODE"].map(tf.ACTION_FILL_HEX).tolist()})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Synthèse", "Détail"]
    ws = wb["Détail"]
    assert ws["A1"].font.bold and ws["A1"].fill.start_color.rgb.endswith("002157")
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:D{len(df) + 1}"
    ligne = 2 + int(df["PSD"].notna().to_numpy().argmax())  # première ligne avec une PSD
    assert ws[f"C{ligne}"].number_format.lower() == "dd/mm/yyyy"
    assert ws[f"C{ligne}"].value.date() == df["PSD"].dropna().iloc[0].date()
    assert ws["A2"].fill.start_color.rgb.endswith(tf.ACTION_FILL_HEX[df["ACTION_CODE"].iloc[0]])


def test_nom_export():
    assert sio.nom_export("stock_chantier", date(2026, 10, 22)) == "Supply360_stock_chantier_20261022.xlsx"


# --------------------------------------------------------------------------
# Reconnaissance des fichiers par leur nom (import groupé d'un dossier)
# --------------------------------------------------------------------------
def test_reconnaitre_fichier():
    attendu = {
        "F-XKAA - 151126 - LOGOPS.xlsx": "besoins",
        "F-XKBB - 030227 - ASREQ.xlsm": "asreq",
        "Etat_stock_20261020.xlsx": "stock",
        "État du Stock.xlsx": "stock",
        "Commandes_SAP_20261020.xlsx": "commandes",
        "References_textile.xlsx": "references",
        "Pièces par avion.csv": "pieces",
        "export (3).xlsx": None,
    }
    assert {n: sio.reconnaitre_fichier(n) for n in attendu} == attendu


def test_chantier_depuis_nom():
    assert sio.chantier_depuis_nom("F-XKAA - 151126 - LOGOPS.xlsx") == ("F-XKAA", date(2026, 11, 15))
    assert sio.chantier_depuis_nom("f-xkbb-030227 asreq.xlsm") == ("F-XKBB", date(2027, 2, 3))
    assert sio.chantier_depuis_nom("F-XKAA - 321326 - LOGOPS.xlsx") is None      # date invalide
    assert sio.chantier_depuis_nom("Etat_stock_20261020.xlsx") is None


def test_detecter_source_par_contenu():
    """Le contenu prime sur le nom : un fichier mal nommé est quand même reconnu."""
    for fichier, source in [("etat_stock.xlsx", "stock"), ("commandes_sap.xlsx", "commandes"), ("besoins_chantier.xlsx", "besoins"),
                            ("etude_asreq.xlsm", "asreq"), ("references_textile.xlsx", "references"), ("pieces_par_avion.csv", "pieces")]:
        data = (FIX / fichier).read_bytes()
        assert sio.detecter_source(data, "export (3)." + fichier.rsplit(".", 1)[1]) == source
    assert sio.detecter_source(b"pas un classeur", "x.xlsx") is None
    assert sio.identifier_fichier(b"pas un classeur", "R2/F-XKAA - 151126/Etat_stock.xlsx") == "stock"   # repli sur le nom


def test_dossiers_regroupement_et_choix():
    from core import dossiers as d
    chemins = ["R2/F-XKAA - 151126/a.xlsx", "R2/F-XKBB - 030227/b.xlsx", "R2/lisez-moi.xlsx", "R2/F-AAAA - 010126/c.xlsx"]
    g = d.regrouper(chemins)
    assert d.chantiers_tries(g) == ["F-AAAA - 010126", "F-XKAA - 151126", "F-XKBB - 030227"]
    assert g[d.HORS_CHANTIER] == ["R2/lisez-moi.xlsx"]
    assert d.date_dans_nom("R2/F-XKAA - 151126/stock 20261020.xlsx") == date(2026, 10, 20)
    assert d.date_dans_nom("stock 15.10.26.xlsx") == date(2026, 10, 15)
    par_source, inconnus = d.repartir({"x/stock 20250101.xlsx": "stock", "x/stock 20261020.xlsx": "stock", "x/photo.xlsx": None})
    assert par_source["stock"][0] == "x/stock 20261020.xlsx" and inconnus == ["x/photo.xlsx"]


def test_regles_de_nommage_textile():
    """Dossier Textile : le nom fait foi (« Résultat… », textile_master_export, textile_all_parts), sinon le contenu."""
    from core import dossiers as d
    regles = {"stock": r"^resultat", "references": r"textile_master_export", "pieces": r"textile_all_parts"}
    sources = {"T/Résultat 20261020.xlsx": "stock", "T/autre stock.xlsx": "stock", "T/textile_master_export.xlsx": "references",
               "T/textile_all_parts.csv": None, "T/ME2N.xlsx": "commandes"}
    r = d.appliquer_noms(sources, regles)
    assert r["T/Résultat 20261020.xlsx"] == "stock" and r["T/autre stock.xlsx"] is None
    assert r["T/textile_all_parts.csv"] == "pieces" and r["T/ME2N.xlsx"] == "commandes"
    assert d.appliquer_noms({"T/stock.xlsx": "stock"}, regles) == {"T/stock.xlsx": "stock"}   # aucun « Résultat » : contenu


def test_tracking_level_lu_dans_etat_du_stock():
    """TRACKING LEVEL → TL « 0 » / « 1 », quelle que soit l'écriture ; absent → vide."""
    df = pd.DataFrame({"PNR": ["A1", "B2", "C3", "D4"], "QTE SUR ÉTAGÈRE": [1, 2, 3, 4],
                       "TRACKING LEVEL": ["TL0", "1", "TL 1", None]})
    lignes, _ = sio.lire_etat_stock(sio.exporter_excel({"Stock": df}))
    assert lignes.set_index("PN")["TL"].to_dict() == {"A1": "0", "B2": "1", "C3": "1", "D4": np.nan} or \
        lignes.set_index("PN")["TL"].fillna("vide").to_dict() == {"A1": "0", "B2": "1", "C3": "1", "D4": "vide"}
    agg = tf.agreger_stock(lignes)
    assert agg.set_index("PN").loc["B2", "TL"] == "1"
    sans, _ = sio.lire_etat_stock(sio.exporter_excel({"Stock": df.drop(columns="TRACKING LEVEL")}))
    assert sans["TL"].isna().all()


def test_apercu_rapide_des_en_tetes(monkeypatch):
    """La reconnaissance par aperçu du XML (rapide) donne le même résultat que la lecture complète des feuilles."""
    for f in sorted(FIX.glob("*.xls*")):
        data = f.read_bytes()
        rapide = sio.detecter_source(data, "x.xlsx")
        xl = sio.ouvrir_excel(data, "x")
        feuille_rapide = sio.choisir_feuille(xl, ["pnr", "qte sur etagere"])
        with monkeypatch.context() as m:
            m.setattr(sio, "apercu_xlsx", lambda *a, **k: None)
            assert sio.detecter_source(data, "x.xlsx") == rapide, f.name
            assert sio.choisir_feuille(sio.ouvrir_excel(data, "x"), ["pnr", "qte sur etagere"]) == feuille_rapide, f.name
    lignes = dict(sio.apercu_xlsx((FIX / "commandes_sap.xlsx").read_bytes()))["Toutes les commandes"]
    assert lignes[3][:2] == ["PN", "Commande"]          # en-tête en 4e ligne, après les lignes de titre


def _stock_variante(tmp_path, decalage=0, prefixe=False, textes_ailleurs=False):
    """État du stock fictif écrit « autrement qu'Excel » : lignes de titre au-dessus de l'en-tête, balises préfixées
    (x:row…), textes partagés rangés sous un autre nom. Ces variantes empêchaient de reconnaître un export réel."""
    import io as _io
    import re as _re
    import zipfile as _zip

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export"
    for i in range(decalage):
        ws.append([f"Titre {i}"])
    ws.append(["PNR", "DESCRIPTION", "QTE SUR ÉTAGÈRE", "QTE EN CDE", "CCN"])
    ws.append(["FX-001", "Housse", 4, 2, "7101"])
    tampon = _io.BytesIO()
    wb.save(tampon)
    entree, sortie = _zip.ZipFile(_io.BytesIO(tampon.getvalue())), _io.BytesIO()
    with _zip.ZipFile(sortie, "w", _zip.ZIP_DEFLATED) as z:
        for info in entree.infolist():
            nom, contenu = info.filename, entree.read(info).decode("utf-8")
            if prefixe and nom.endswith((".xml",)) and nom.startswith("xl/") and "theme" not in nom and "styles" not in nom:
                contenu = _re.sub(r"<(/?)(?!\?)([A-Za-z]+)([ >/])", r"<\1x:\2\3", contenu)
                contenu = _re.sub(r'xmlns="([^"]+)"', r'xmlns:x="\1"', contenu, count=1)
                contenu = "\ufeff" + contenu  # BOM en tête, comme l'export réel
            if prefixe and nom == "xl/_rels/workbook.xml.rels":  # cibles absolues, identifiants non standard
                contenu = contenu.replace('Target="worksheets/', 'Target="/xl/worksheets/')
            if prefixe and nom in ("xl/workbook.xml", "xl/_rels/workbook.xml.rels"):
                contenu = contenu.replace('"rId1"', '"feuilleexport"')
            if textes_ailleurs:
                contenu = contenu.replace("sharedStrings.xml", "Textes/Partages.xml")
                nom = "xl/Textes/Partages.xml" if nom == "xl/sharedStrings.xml" else nom
            z.writestr(nom, contenu)
    f = tmp_path / "Résultat - HZUA.xlsx"
    f.write_bytes(sortie.getvalue())
    return f.read_bytes()


@pytest.mark.parametrize("variante", [dict(decalage=20), dict(prefixe=True), dict(textes_ailleurs=True),
                                      dict(decalage=18, prefixe=True, textes_ailleurs=True)])
def test_etat_du_stock_reconnu_quelle_que_soit_son_ecriture(tmp_path, variante):
    data = _stock_variante(tmp_path, **variante)
    assert sio.detecter_source(data, "export.xlsx") == "stock"
    stock = sio.lire_etat_stock(data)[0]
    assert list(stock["PN"]) == ["FX-001"]


def test_resultat_reconnu_par_son_nom():
    assert sio.reconnaitre_fichier("Résultat - HZUA.xlsx") == "stock"
    assert sio.identifier_fichier(b"pas un classeur", "R2/F-HZUA - 110127/Résultat - HZUA.xlsx") == "stock"


def test_feuille_annexe_ne_masque_pas_l_etat_du_stock():
    # 1re feuille qui ressemble à des Références textile (PN, Équipement), données de stock en 2e feuille
    wb = Workbook()
    wb.active.title = "Légende"
    wb.active.append(["PN", "Équipement", "Commentaire"])
    ws = wb.create_sheet("Données")
    ws.append(["PNR", "QTE SUR ÉTAGÈRE", "EQUIPEMENT"])
    ws.append(["FX-001", 3, "Siège"])
    tampon = io.BytesIO()
    wb.save(tampon)
    assert sio.detecter_source(tampon.getvalue(), "Résultat - HZUA.xlsx") == "stock"


def test_dossier_r2_resultat_fait_foi():
    from core import dossiers
    from modules import stock_chantier

    sources = {"R2/F-HZUA - 110127/Résultat - HZUA.xlsx": "references", "R2/F-HZUA - 110127/vieux stock.xlsx": "stock",
               "R2/F-HZUA - 110127/LOGOPS.xlsx": "besoins"}
    par_source, inconnus = dossiers.repartir(dossiers.appliquer_noms(sources, stock_chantier.NOMS_DOSSIER))
    assert par_source["stock"] == ["R2/F-HZUA - 110127/Résultat - HZUA.xlsx"] and par_source["besoins"]
    assert inconnus == ["R2/F-HZUA - 110127/vieux stock.xlsx"]


def test_stock_chantier_colonnes_rrs_sar_pr_code_agent(sources, croise):
    from modules import stock_chantier

    stock = sources[1].set_index("PN")
    assert stock.loc["FX-1003", "QTE_SAR"] == 1 and stock.loc["FX-1002", "QTE_PR"] == 2
    m = croise.set_index("PN")
    # contact = CODE AGT APPRO de l'État du stock (plus le code acheteur des commandes)
    assert "(contact: AG01)" in m.loc["FX-1002", "ACTION"]
    assert m.loc["FX-1002", "Qty_a_commander"] == 4   # besoin 5 (2 + 3, deux TRFX) − stock 1
    t = stock_chantier.tableau(croise)
    assert list(t.columns)[:19] == ["Statut", "PN", "CCN", "RRS", "À commander", "Besoin", "Stock", "En cmd", "SAR", "PR",
                                    "N° commande", "PSD", "Fournisseur", "Code fournisseur", "Action", "Code agent",
                                    "Gestion", "TL", "Date chantier"]
    assert "Acheteur" not in t.columns
